"""
محاكاة دورة كاملة (بروفة) — تُنشئ دورة «محاكاة» منفصلة ببيانات وهمية معلّمة «تجريبي» وتمرّ بكل مراحل الجائزة،
وتفحص كل خطوة (نجح/فشل) وتكتب تقريراً. لا تلمس الدورة الحالية ولا ترسل أي بريد ولا تنبّه الإدارة.

    python manage.py simulate_cycle            # تشغيل
    python manage.py simulate_cycle --clear    # مسح كل بيانات التجربة

حسابات المحاكاة: sim-…@sim.iaj.invalid بكلمة مرور واحدة تظهر في التقرير — يمكنكم الدخول بها لتروا ما تراه
المدرسة والمحكّم. كل شيء يُحذف بزر «مسح بيانات التجربة» في /manage/tools/.
"""
import io
import random
import secrets
import traceback
from datetime import timedelta
from decimal import Decimal

from django.conf import settings
from django.contrib.auth.models import User
from django.test import Client
from django.utils import timezone

from .models import (SIM_DOMAIN, Appeal, Assignment, AwardCycle, Certificate, Criterion, DecisionTemplate, EmailLog,
                     Governorate, ImpactReport, Invitation, Notification, Profile, Score, SimulationRun, StatusLog,
                     Submission, Track, Winner)
from .notify import Quiet
from . import workflow, reminders


class Abort(Exception):
    pass


class Sim:
    def __init__(self, by=None, n_schools=10, n_judges=6, seed=None):
        self.by = by
        self.n_schools = max(8, n_schools)
        self.n_judges = max(4, n_judges)
        self.rnd = random.Random(seed if seed is not None else secrets.randbits(32))
        self.lines = []
        self.passed = self.failed = 0
        self.password = 'Sim-' + secrets.token_hex(4)
        self.stamp = timezone.localtime().strftime('%Y-%m-%d %H:%M')
        self.cycle = None

    # ---------- أدوات ----------
    def log(self, text):
        self.lines.append(text)

    def check(self, name, ok, detail=''):
        if ok:
            self.passed += 1
            self.log(f'  ✓ {name}' + (f' — {detail}' if detail else ''))
        else:
            self.failed += 1
            self.log(f'  ✗ {name}' + (f' — {detail}' if detail else ''))
        return ok

    def must(self, name, ok, detail=''):
        if not self.check(name, ok, detail):
            raise Abort(name)

    def client(self, user):
        host = next((h for h in settings.ALLOWED_HOSTS if h and '*' not in h and not h.startswith('.')), 'localhost')
        c = Client(HTTP_HOST=host)
        c.force_login(user)
        return c

    def get(self, user, url):
        return self.client(user).get(url, secure=True)

    def post(self, user, url, data, files=None):
        return self.client(user).post(url, data, secure=True)

    def _user(self, kind, i, name, role, **prof):
        email = f'sim-{kind}-{i:02d}-{secrets.token_hex(2)}@{SIM_DOMAIN}'
        u = User.objects.create_user(username=email, email=email, password=self.password, first_name=name[:150])
        Profile.objects.create(user=u, role=role, contact_person=name, phone=f'07900000{i:02d}', is_simulation=True, **prof)
        return u

    # ---------- التشغيل ----------
    def run(self):
        steps = [
            ('١. تجهيز دورة المحاكاة والمعايير', self.s_setup),
            ('٢. حسابات المدارس والمحكّمين', self.s_accounts),
            ('٣. التقديم: المسودة والإقرارات والإرسال والإيصال', self.s_submit),
            ('٤. التذكيرات قبل الإغلاق', self.s_reminders_school),
            ('٥. التدقيق: قبول للتدقيق / طلب استكمال / تعليق / رفض / سحب', self.s_screening),
            ('٦. التحويل للتحكيم وتعارض المصالح والتوزيع المتوازن', self.s_assign),
            ('٧. تذكير المحكّمين المتأخرين', self.s_reminders_judges),
            ('٨. التقييم وكشف التباين والمحكّم الإضافي', self.s_scoring),
            ('٩. قفل الدرجات بعد موعد التحكيم', self.s_lock),
            ('١٠. مراحل التحكيم: القائمة القصيرة ثم النهائي', self.s_stages),
            ('١١. الترتيب وكسر التعادل', self.s_ranking),
            ('١٢. الفائزون واعتماد اللجنة ثم النشر', self.s_results),
            ('١٣. تقرير التغذية الراجعة للمدارس', self.s_feedback),
            ('١٤. الاعتراضات', self.s_appeals),
            ('١٥. الشهادات والتحقق برمز QR', self.s_certificates),
            ('١٦. دعوات الحفل والرد (RSVP)', self.s_invitations),
            ('١٧. متابعة أثر الفائزين', self.s_impact),
            ('١٨. تصدير Excel للدورة', self.s_export),
            ('١٩. صفحات الإدارة تعمل', self.s_pages),
            ('٢٠. الخصوصية: لا تظهر بيانات المحاكاة للزوار ولا يُرسل بريد', self.s_privacy),
        ]
        self.emails_before = EmailLog.objects.count()
        self.log(f'محاكاة دورة كاملة — {self.stamp}')
        with Quiet():
            for title, fn in steps:
                self.log('')
                self.log(title)
                try:
                    fn()
                except Abort as e:
                    self.log(f'  ⛔ توقفت المحاكاة: فشل أساسي في «{e}»')
                    break
                except Exception as e:   # noqa: BLE001
                    self.failed += 1
                    self.log(f'  ✗ خطأ غير متوقع: {type(e).__name__}: {e}')
                    self.log('    ' + traceback.format_exc().strip().splitlines()[-3][:300])
                    break
        self.log('')
        self.log(f'النتيجة: {self.passed} نجح · {self.failed} فشل')
        self.log(f'حسابات المحاكاة: كلمة المرور {self.password}  (أمثلة: {getattr(self, "example_school", "")} · {getattr(self, "example_judge", "")})')
        self.log('لمسح كل بيانات المحاكاة: /manage/tools/ ← «مسح بيانات التجربة»  أو  simulate_clear.bat')
        run = SimulationRun.objects.create(by=self.by if self.by and self.by.is_authenticated else None, cycle=self.cycle,
                                           passed=self.passed, failed=self.failed, report='\n'.join(self.lines),
                                           password=self.password)
        from .audit import audit
        audit(self.by, 'محاكاة دورة', str(self.cycle or ''), f'{self.passed} نجح / {self.failed} فشل')
        return run

    # ---------- الخطوات ----------
    def s_setup(self):
        tracks = list(Track.objects.filter(is_active=True).select_related('field')[:3])
        self.must('توجد مسارات مفعّلة في الموقع', bool(tracks), f'{len(tracks)} مسار')
        now = timezone.now()
        self.cycle = AwardCycle.objects.create(
            name=f'محاكاة {self.stamp}', short_name='محاكاة', year=now.year, is_current=False, is_simulation=True,
            opens_at=now - timedelta(days=20), closes_at=now + timedelta(days=2), judges_per_submission=2,
            revision_days=5, blind_judging=True, discrepancy_points=20, tie_break='criterion', appeal_days=7,
            judging_closes_at=now + timedelta(days=2), ceremony_date=now + timedelta(days=30),
            ceremony_place='قاعة المحاكاة', show_card=False, impact_after_days=180)
        self.cycle.tracks.set(tracks)
        self.tracks = tracks
        # المعيار الأول وزنه الأكبر (لاختبار كسر التعادل)
        self.crits = [Criterion.objects.create(cycle=self.cycle, name=n, max_score=10, weight=Decimal(w), order=i)
                      for i, (n, w) in enumerate([('الأصالة والابتكار', '5'), ('المنهجية', '1'), ('الأثر', '1'), ('العرض', '1')])]
        self.check('أُنشئت دورة محاكاة منفصلة (ليست الحالية)', not self.cycle.is_current and AwardCycle.current() != self.cycle,
                   self.cycle.name)
        self.check('معايير الدورة', len(Criterion.for_submission(Submission(cycle=self.cycle, track=tracks[0]))) >= 4,
                   '، '.join(c.name for c in self.crits))

    def s_accounts(self):
        gov = Governorate.objects.first()
        self.manager = self._user('manager', 1, 'مدير المحاكاة', 'manager')
        self.judges = [self._user('judge', i + 1, f'محكّم محاكاة {i + 1}', 'judge', specialty='محاكاة')
                       for i in range(self.n_judges)]
        self.schools = [self._user('school', i + 1, f'منسق مدرسة {i + 1}', 'school',
                                   school_name=f'مدرسة المحاكاة {i + 1}', governorate=gov)
                        for i in range(self.n_schools)]
        self.example_school, self.example_judge = self.schools[0].email, self.judges[0].email
        from .portal_views import is_manager
        self.check('حسابات المدارس', len(self.schools) == self.n_schools, f'{len(self.schools)} مدرسة')
        self.check('حسابات المحكّمين', len(self.judges) == self.n_judges, f'{len(self.judges)} محكّم')
        self.must('حساب مدير المحاكاة يدخل لوحة الإدارة', is_manager(self.manager))

    def s_submit(self):
        self.subs = []
        for i, u in enumerate(self.schools):
            track = self.tracks[0] if i in (6, 7) else self.tracks[i % len(self.tracks)]
            p = u.profile
            s = Submission.objects.create(
                cycle=self.cycle, owner=u, school_name=p.school_name, contact_person=p.contact_person, email=u.email,
                phone=p.phone, field=track.field, track=track, project_title=f'مشروع محاكاة رقم {i + 1}',
                abstract='ملخص تجريبي للمشروع لأغراض المحاكاة.', team_members='طالب أ\nطالب ب', supervisor='معلم المحاكاة',
                status='draft', is_simulation=True)
            StatusLog.objects.create(submission=s, new_status='draft', by=u)
            self.subs.append(s)
        self.draft = self.subs[-1]   # آخر مدرسة تبقى مسودة (لاختبار التذكير)
        for s in self.subs[:-1]:
            s.declared_truth = s.declared_original = s.declared_consent = True
            s.declared_at = timezone.now()
            s.document.name = 'private/submissions/sim/research.pdf'
            s.save()
            workflow.change_status(s, 'submitted', by=s.owner)
        sent = [s for s in self.subs if s.status == 'submitted']
        self.check('الطلبات أُرسلت ولها رقم وتاريخ إرسال', all(x.ref and x.sent_at for x in sent), f'{len(sent)} طلب')
        self.check('الإقرارات محفوظة', all(x.declared_at for x in sent))
        self.check('إشعار «تم استلام طلبك» لكل مدرسة',
                   Notification.objects.filter(user__in=[x.owner for x in sent], title__contains='تم استلام').count() == len(sent))
        r = self.get(self.schools[0], f'/portal/submissions/{sent[0].pk}/receipt/')
        self.check('إيصال الاستلام يُفتح للمدرسة', r.status_code == 200 and sent[0].ref in r.content.decode(), f'HTTP {r.status_code}')
        r = self.get(self.schools[1], f'/portal/submissions/{sent[0].pk}/receipt/')
        self.check('مدرسة أخرى لا ترى إيصال غيرها', r.status_code == 404, f'HTTP {r.status_code}')

    def s_reminders_school(self):
        n = reminders.draft_reminders(self.cycle, force=True)
        self.check('تذكير المسودات غير المرسلة', n == 1 and Notification.objects.filter(user=self.draft.owner, title__contains='لم يُرسل').exists(),
                   f'{n} تذكير')
        self.check('لا يتكرر التذكير نفسه', reminders.draft_reminders(self.cycle, force=True) == 0)

    def s_screening(self):
        s = self.subs
        for x in s[:-1]:
            workflow.change_status(x, 'screening', by=self.manager)
        self.check('قيد التدقيق', all(x.status == 'screening' for x in s[:-1]))
        tpl = DecisionTemplate.objects.filter(status='revision').first()
        note = tpl.body if tpl else 'يرجى إرفاق صور التنفيذ.'
        workflow.change_status(s[1], 'revision', by=self.manager, note=note)
        s[1].refresh_from_db()
        self.check('طلب استكمال بمهلة وسبب من القوالب', s[1].revision_deadline is not None and s[1].can_edit and s[1].school_note == note,
                   f'آخر موعد {timezone.localtime(s[1].revision_deadline):%Y-%m-%d}')
        self.check('تذكير قرب انتهاء مهلة الاستكمال', reminders.revision_reminders(force=True) >= 1)
        workflow.change_status(s[1], 'submitted', by=s[1].owner)
        workflow.change_status(s[1], 'screening', by=self.manager)
        self.check('المدرسة أعادت الإرسال بعد التعديل', s[1].logs.filter(old_status='revision', new_status='submitted').exists())
        workflow.change_status(s[2], 'on_hold', by=self.manager, note='بانتظار التحقق من البيانات')
        s[2].refresh_from_db()
        self.check('تعليق مع سبب يصل للمدرسة', s[2].status == 'on_hold' and 'التحقق' in s[2].school_note and
                   Notification.objects.filter(user=s[2].owner, title__contains='معلّق').exists())
        workflow.change_status(s[2], 'screening', by=self.manager)
        workflow.change_status(s[3], 'rejected', by=self.manager, note='خارج نطاق المسار')
        s[3].refresh_from_db()
        self.check('رفض: لا يظهر للمدرسة قبل نشر النتائج', s[3].results_hidden and s[3].public_status == 'judging' and
                   not Notification.objects.filter(user=s[3].owner, title__contains='نتيجة').exists())
        workflow.change_status(s[4], 'withdrawn', by=s[4].owner, notify_school=False)
        self.check('سحب الطلب من المدرسة', s[4].status == 'withdrawn')
        self.judging = [x for x in s[:-1] if x.status == 'screening']
        for x in self.judging:
            workflow.change_status(x, 'judging', by=self.manager)
        self.check('تحويل للتحكيم', all(x.status == 'judging' for x in self.judging), f'{len(self.judging)} طلب')

    def s_assign(self):
        j0, sub0 = self.judges[0], self.judging[0]
        workflow.declare_conflict(j0, None, sub0.school_name, 'قرابة مع المشرف (محاكاة)', by=j0)
        n = workflow.auto_assign(self.judging)
        self.check('توزيع تلقائي', n == 2 * len(self.judging), f'{n} إسناد')
        self.check('المحكّم صاحب التعارض لم يُسند له طلب تلك المدرسة',
                   not Assignment.objects.filter(judge=j0, submission=sub0).exists())
        loads = [Assignment.objects.filter(judge=j).count() for j in self.judges]
        self.check('توزيع متوازن للحمل', max(loads) - min(loads) <= 2, f'الحمل: {loads}')
        real = Assignment.objects.filter(submission__in=self.judging, judge__profile__is_simulation=False).count()
        self.check('لم يُسند أي طلب تجريبي لمحكّم حقيقي', real == 0)
        # تعارض يعلنه المحكّم من صفحة التقييم نفسها
        a = Assignment.objects.filter(submission=self.judging[1]).first()
        r = self.post(a.judge, f'/judge/{a.pk}/', {'action': 'coi_conflict', 'reason': 'محاكاة', 'scope': 'submission'})
        self.check('إعلان التعارض من صفحة المحكّم يلغي الإسناد', r.status_code == 302 and not Assignment.objects.filter(pk=a.pk).exists())
        workflow.auto_assign([self.judging[1]])
        self.check('إسناد بديل تلقائي', Assignment.objects.filter(submission=self.judging[1]).count() == 2 and
                   not Assignment.objects.filter(submission=self.judging[1], judge=a.judge).exists())
        a = Assignment.objects.filter(submission=self.judging[2]).first()
        r = self.post(a.judge, f'/judge/{a.pk}/', {'action': 'finish', **{f'c{c.pk}': 5 for c in self.crits}})
        a.refresh_from_db()
        self.check('لا تقييم قبل إقرار عدم التعارض', not a.is_done)
        Assignment.objects.filter(submission__in=self.judging).update(coi_confirmed_at=timezone.now())

    def s_reminders_judges(self):
        n = reminders.judge_reminders(self.cycle, force=True)
        self.check('تذكير المحكّمين بالطلبات المتبقية', n >= 1 and Notification.objects.filter(
            user__in=self.judges, title__contains='بانتظار تقييمك').exists(), f'{n} محكّم')

    def _score(self, a, vals, rec='yes', via_web=False):
        crits = list(Criterion.for_submission(a.submission))
        if via_web:
            data = {'action': 'finish', 'comment': 'تقييم محاكاة', 'recommendation': rec}
            for c in crits:
                data[f'c{c.pk}'] = vals.get(c.pk, round(c.max_score * 0.7))
            return self.post(a.judge, f'/judge/{a.pk}/', data)
        for c in crits:
            Score.objects.update_or_create(assignment=a, criterion=c, defaults={'value': vals.get(c.pk, round(c.max_score * 0.7))})
        a.recommendation, a.comment, a.completed_at = rec, 'ملاحظة محاكاة', timezone.now()
        a.save()
        workflow.check_discrepancy(a.submission)

    def s_scoring(self):
        c1, c2, c3, c4 = self.crits
        x_sub, y_sub = self.subs[6], self.subs[7]   # نفس المسار — متساويان في المتوسط (اختبار التعادل)
        disc = self.subs[8]
        self.tie_x, self.tie_y, self.disc = x_sub, y_sub, disc
        first = True
        for s in self.judging:
            assigns = list(Assignment.objects.filter(submission=s).select_related('judge', 'submission__cycle'))
            q = self.rnd.uniform(0.45, 0.9)
            for k, a in enumerate(assigns):
                if s == x_sub:
                    vals, rec = {c1.pk: 8, c2.pk: 6, c3.pk: 6, c4.pk: 6}, 'yes'
                elif s == y_sub:
                    vals, rec = {c1.pk: 6, c2.pk: 10, c3.pk: 10, c4.pk: 8}, 'yes'
                elif s == disc:
                    v = 9 if k == 0 else 3
                    vals, rec = {c.pk: v for c in self.crits}, 'strong' if k == 0 else 'no'
                else:
                    vals, rec = {c.pk: max(0, min(10, round(q * 10 + self.rnd.uniform(-0.5, 0.5)))) for c in self.crits}, \
                        self.rnd.choice(['strong', 'yes', 'maybe'])
                if first:
                    r = self._score(a, vals, rec, via_web=True)
                    a.refresh_from_db()
                    self.check('المحكّم يقيّم من صفحة التقييم (الموقع)', a.is_done, f'HTTP {r.status_code} — {a.total}/100')
                    first = False
                else:
                    self._score(a, vals, rec)
        disc.refresh_from_db()
        n = Assignment.objects.filter(submission=disc).count()
        self.check('كشف تباين الدرجات', disc.has_discrepancy, f'الفرق {disc.score_spread} ≥ {self.cycle.discrepancy_points}')
        self.check('إسناد محكّم إضافي تلقائياً عند التباين', n == 3, f'{n} محكّمين')
        extra = Assignment.objects.filter(submission=disc, completed_at__isnull=True).first()
        if extra:
            extra.coi_confirmed_at = timezone.now()
            extra.save()
            self._score(extra, {c.pk: 6 for c in self.crits}, 'maybe')
        self.check('لا يُضاف محكّم رابع (مرة واحدة فقط)', Assignment.objects.filter(submission=disc).count() == 3)
        pending = Assignment.objects.filter(submission__in=self.judging, completed_at__isnull=True).count()
        self.check('اكتمل التحكيم الأولي', pending == 0, f'{pending} متبقٍ')

    def s_lock(self):
        s = self.judging[0]
        a = Assignment.objects.filter(submission=s, completed_at__isnull=False).first()
        before = a.total
        self.cycle.judging_closes_at = timezone.now() - timedelta(minutes=1)
        self.cycle.save()
        self.post(a.judge, f'/judge/{a.pk}/', {'action': 'finish', **{f'c{c.pk}': 0 for c in self.crits}})
        a = Assignment.objects.get(pk=a.pk)
        self.check('الدرجات مقفلة بعد الموعد', a.total == before, f'{before} لم تتغير')
        r = self.get(a.judge, f'/judge/{a.pk}/')
        self.check('المحكّم يرى رسالة القفل', 'مقفلة' in r.content.decode())
        self.cycle.judging_closes_at = timezone.now() + timedelta(days=2)   # تمديد
        self.cycle.save()

    def s_stages(self):
        xy = {self.tie_x.pk, self.tie_y.pk}
        ranked = sorted((Submission.objects.get(pk=s.pk) for s in self.judging if s.pk not in xy),
                        key=lambda s: -(s.avg_score or 0))
        top3 = ranked[:3]
        short = top3 + [Submission.objects.get(pk=self.tie_x.pk), Submission.objects.get(pk=self.tie_y.pk)]
        n = workflow.set_stage(short, 'shortlist', by=self.manager)
        self.check('نقل إلى القائمة القصيرة', n == len(short), f'{n} طلب')
        made = workflow.auto_assign(short)
        self.check('محكّمون جدد للقائمة القصيرة', made == 2 * len(short), f'{made} إسناد')
        prev = sum(1 for s in short for a in s.assignments.all()
                   if a.stage == 'shortlist' and s.assignments.filter(stage='initial', judge=a.judge).exists())
        self.check('الأولوية لمحكّمين لم يقيّموا الطلب سابقاً', prev <= len(short), f'{prev} تكرار')
        c1, c2, c3, c4 = self.crits
        for s in short:
            for a in Assignment.objects.filter(submission=s, stage='shortlist'):
                a.coi_confirmed_at = timezone.now()
                a.save()
                if s.pk == self.tie_x.pk:
                    vals = {c1.pk: 8, c2.pk: 6, c3.pk: 6, c4.pk: 6}
                elif s.pk == self.tie_y.pk:
                    vals = {c1.pk: 6, c2.pk: 10, c3.pk: 10, c4.pk: 8}
                else:
                    q = self.rnd.uniform(0.6, 0.95)
                    vals = {c.pk: round(q * 10) for c in self.crits}
                self._score(a, vals, 'yes')
        s0 = Submission.objects.get(pk=short[0].pk)
        st = [a.total for a in s0.assignments.all() if a.stage == 'shortlist']
        self.check('المتوسط يُحسب من المرحلة الحالية فقط', s0.avg_score == round(sum(st) / len(st), 1), f'{s0.avg_score}')
        final = [Submission.objects.get(pk=s.pk) for s in top3]
        workflow.set_stage(final, 'final', by=self.manager)
        made = workflow.auto_assign(final, per_submission=3)
        for s in final:
            for a in Assignment.objects.filter(submission=s, stage='final'):
                a.coi_confirmed_at = timezone.now()
                a.save()
                self._score(a, {c.pk: self.rnd.choice([7, 8, 9]) for c in self.crits}, 'strong')
        self.check('التحكيم النهائي لأفضل 3', made == 9 and all(Submission.objects.get(pk=s.pk).judging_stage == 'final' for s in final),
                   f'{made} إسناد')
        self.short, self.final = short, final

    def s_ranking(self):
        x = Submission.objects.get(pk=self.tie_x.pk)
        y = Submission.objects.get(pk=self.tie_y.pk)
        dbg = ' | '.join(f"{s.ref}:{s.judging_stage}:" + ','.join(f"{a.stage}/{a.total}" for a in s.assignments.all()) for s in (x, y))
        self.check('طلبان بنفس المتوسط (تعادل)', x.avg_score == y.avg_score, f'{x.avg_score} = {y.avg_score}' + ('' if x.avg_score == y.avg_score else ' ' + dbg))
        rows = workflow.rank_rows([y, x], 'criterion')
        self.check('كسر التعادل بالمعيار الأكبر وزناً', rows[0]['s'].pk == x.pk and rows[0]['tie'],
                   f'{rows[0]["s"].ref} قبل {rows[1]["s"].ref}')
        rows = workflow.rank_rows([x, y], 'min')
        self.check('قاعدة بديلة (أعلى أدنى درجة) تعمل', len(rows) == 2)

    def s_results(self):
        final = [Submission.objects.get(pk=s.pk) for s in self.final]
        final.sort(key=lambda s: -(s.avg_score or 0))
        for rank, s in enumerate(final, 1):
            workflow.make_winner(s, rank=rank, by=self.manager, note='مبارك')
        others = Submission.objects.filter(pk__in=[s.pk for s in self.judging]).exclude(status='winner')
        accepted = [s for s in others if s.judging_stage == 'shortlist']
        for s in others:
            workflow.change_status(s, 'accepted' if s in accepted else 'rejected', by=self.manager,
                                   note='' if s in accepted else 'لم يتأهل بعد التحكيم')
        self.check('اعتماد الفائزين', Winner.objects.filter(cycle=self.cycle).count() == 3)
        try:
            workflow.publish_results(self.cycle, self.manager)
            self.check('النشر ممنوع قبل اعتماد اللجنة', False)
        except workflow.NotApproved:
            self.check('النشر ممنوع قبل اعتماد اللجنة', True)
        workflow.approve_results(self.cycle, self.manager, 'محضر اجتماع اللجنة (محاكاة)')
        self.cycle.refresh_from_db()
        self.check('تسجيل اعتماد اللجنة (من ومتى)', self.cycle.results_approved_by_id == self.manager.pk)
        workflow.make_winner(final[0], rank=1, by=self.manager)   # تعديل بعد الاعتماد
        self.cycle.refresh_from_db()
        self.check('أي تعديل على الفائزين بعد الاعتماد يلغيه', not self.cycle.is_approved)
        workflow.approve_results(self.cycle, self.manager, 'اعتماد نهائي')
        before = Notification.objects.filter(user__in=self.schools).count()
        n = workflow.publish_results(self.cycle, self.manager)
        finals = self.cycle.submissions.filter(status__in=('accepted', 'rejected', 'winner')).count()
        self.cycle.refresh_from_db()
        self.check('نشر النتائج وإبلاغ كل المدارس', n == finals and Notification.objects.filter(user__in=self.schools).count() - before == finals,
                   f'{n} إشعار')
        w = Submission.objects.get(pk=final[0].pk)
        self.check('المدرسة ترى «فائز» بعد النشر فقط', not w.results_hidden and w.public_status == 'winner')
        self.winners = final

    def s_feedback(self):
        s = Submission.objects.get(pk=self.subs[5].pk)
        fb = workflow.feedback(s)
        self.check('تقرير لكل معيار', len(fb['rows']) >= 4, '، '.join(f"{r['name']} {r['pct']}%" for r in fb['rows'][:4]))
        r = self.get(s.owner, s.get_absolute_url())
        html = r.content.decode()
        self.check('يظهر للمدرسة في صفحة الطلب', 'تقرير التغذية الراجعة' in html)
        self.check('بدون أسماء المحكّمين', not any(j.first_name in html for j in self.judges))

    def s_appeals(self):
        rej = self.cycle.submissions.filter(status='rejected').exclude(pk=self.subs[3].pk).first() or \
            self.cycle.submissions.filter(status='rejected').first()
        self.must('يوجد طلب مرفوض للاعتراض', rej is not None)
        self.check('الاعتراض متاح ضمن المهلة', workflow.can_appeal(rej), f'حتى {timezone.localtime(self.cycle.appeal_until):%Y-%m-%d}')
        r = self.post(rej.owner, f'/portal/submissions/{rej.pk}/appeal/', {'body': 'نرجو إعادة النظر في درجة الأثر لأن المشروع طُبّق فعلاً في المدرسة.'})
        ap = Appeal.objects.filter(submission=rej).first()
        self.check('المدرسة قدّمت اعتراضاً', ap is not None, f'HTTP {r.status_code}')
        self.check('اعتراض واحد فقط', not workflow.can_appeal(rej))
        r = self.post(self.manager, '/manage/appeals/', {'aid': ap.pk, 'status': 'rejected', 'response': 'درست اللجنة الاعتراض والنتيجة ثابتة.'})
        ap.refresh_from_db()
        self.check('رد اللجنة ووصل للمدرسة', ap.status == 'rejected' and Notification.objects.filter(
            user=rej.owner, title__contains='اعتراض').exists(), f'HTTP {r.status_code}')

    def s_certificates(self):
        n = workflow.issue_certificates(self.cycle)
        finals = self.cycle.submissions.filter(status__in=('accepted', 'rejected', 'winner')).count()
        nj = User.objects.filter(assignments__submission__cycle=self.cycle, assignments__completed_at__isnull=False).distinct().count()
        self.check('إصدار الشهادات (فوز/مشاركة/شكر المحكّمين)', n == finals + nj, f'{n} شهادة')
        self.check('لا تتكرر عند الإصدار مرة أخرى', workflow.issue_certificates(self.cycle) == 0)
        c = Certificate.objects.filter(cycle=self.cycle, kind='winner').first()
        r = Client(HTTP_HOST=self.client(self.manager).defaults['HTTP_HOST']).get(f'/verify/{c.code}/', secure=True)
        self.check('صفحة التحقق العامة تؤكد الشهادة', r.status_code == 200 and 'شهادة صحيحة' in r.content.decode(), c.code)
        r = Client(HTTP_HOST=self.client(self.manager).defaults['HTTP_HOST']).get(f'/certificate/{c.code}/', secure=True)
        self.check('الشهادة قابلة للطباعة مع رمز QR', r.status_code == 200 and '<svg' in r.content.decode())
        r = Client(HTTP_HOST=self.client(self.manager).defaults['HTTP_HOST']).get('/verify/XXXXXXXXXX/', secure=True)
        self.check('رمز غير صحيح يُرفض', 'لا توجد شهادة' in r.content.decode())

    def s_invitations(self):
        n = workflow.send_invitations(self.cycle)
        self.check('دعوات للفائزين', n == 3 and Notification.objects.filter(title__contains='دعوة لحضور', user__in=self.schools).count() == 3)
        inv = Invitation.objects.filter(cycle=self.cycle).first()
        host = self.client(self.manager).defaults['HTTP_HOST']
        Client(HTTP_HOST=host).post(f'/rsvp/{inv.code}/', {'rsvp': 'yes', 'guests': '4', 'note': 'المدير والمشرف وطالبان'}, secure=True)
        inv.refresh_from_db()
        self.check('تأكيد الحضور وعدد الحضور', inv.rsvp == 'yes' and inv.guests == 4)

    def s_impact(self):
        self.check('لا يُطلب تقرير الأثر قبل موعده', reminders.impact_reminders(self.cycle) == 0)
        self.cycle.results_published_at = timezone.now() - timedelta(days=self.cycle.impact_after_days + 1)
        self.cycle.save()
        n = reminders.impact_reminders(self.cycle)
        self.check('طلب تقرير متابعة من الفائزين بعد المدة', n == 3, f'{n} مدرسة')
        w = self.winners[0]
        self.post(w.owner, f'/portal/submissions/{w.pk}/impact/', {'body': 'طُبّق المشروع في ثلاث مدارس مجاورة واستفاد منه 400 طالب.', 'beneficiaries': '400'})
        self.check('المدرسة أرسلت تقرير الأثر', ImpactReport.objects.filter(submission=w, beneficiaries=400).exists())
        self.cycle.results_published_at = timezone.now()
        self.cycle.save()

    def s_export(self):
        from openpyxl import load_workbook
        from .export_xlsx import cycle_workbook
        data = cycle_workbook(self.cycle)
        wb = load_workbook(io.BytesIO(data))
        self.check('ملف Excel بكل الأوراق', len(wb.sheetnames) == 8, '، '.join(wb.sheetnames))
        self.check('كل الطلبات في الملف', wb['الطلبات'].max_row - 1 == len(self.subs), f'{wb["الطلبات"].max_row - 1} صف')

    def s_pages(self):
        c = self.cycle.pk
        sub = self.winners[0]
        for url in ('/manage/', f'/manage/submissions/?cycle={c}', f'/manage/submissions/{sub.pk}/', f'/manage/results/?cycle={c}',
                    '/manage/judges/', '/manage/schools/', '/manage/appeals/', f'/manage/post/?cycle={c}', '/manage/tools/',
                    '/manage/audit/'):
            r = self.get(self.manager, url)
            self.check(f'صفحة {url.split("?")[0]}', r.status_code == 200, f'HTTP {r.status_code}')
        r = self.get(self.schools[0], '/manage/')
        self.check('المدرسة لا تدخل لوحة الإدارة', r.status_code == 404, f'HTTP {r.status_code}')

    def s_privacy(self):
        host = self.client(self.manager).defaults['HTTP_HOST']
        anon = Client(HTTP_HOST=host)
        for url in ('/winners/', '/cycles/', '/statistics/'):
            r = anon.get(url, secure=True)
            html = r.content.decode()
            self.check(f'{url} لا يعرض بيانات المحاكاة', r.status_code == 200 and 'مدرسة المحاكاة' not in html and 'محاكاة ' + self.stamp not in html,
                       f'HTTP {r.status_code}')
        sent = EmailLog.objects.filter(pk__gt=0).count() - self.emails_before
        self.check('لم يُرسل أي بريد أثناء المحاكاة', sent == 0, f'{sent} رسالة')

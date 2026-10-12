"""تصدير Excel كامل لدورة: الطلبات، التحكيم التفصيلي، الترتيب، الفائزون، الاعتراضات، الدعوات، المدارس، المحكّمون."""
import io

from django.http import HttpResponse
from django.utils import timezone

from .models import STATUS_CHOICES, Appeal, Assignment, Invitation, Profile, Submission, Winner

STATUS = dict(STATUS_CHOICES)


def _dt(v):
    return timezone.localtime(v).strftime('%Y-%m-%d %H:%M') if v else ''


def _sheet(wb, title, header, rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='0A1632')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(['' if v is None else v for v in r])
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(ws.cell(row=j, column=i).value or ''))
                                       for j in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, max(10, width + 2))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def cycle_workbook(cycle, subs=None):
    from openpyxl import Workbook
    from . import workflow
    wb = Workbook()
    wb.remove(wb.active)
    base = Submission.objects.all() if subs is None else subs
    if cycle and subs is None:
        base = base.filter(cycle=cycle)
    base = (base.select_related('track', 'field', 'cycle', 'owner__profile__governorate', 'owner__profile__directorate')
            .prefetch_related('assignments__scores__criterion', 'assignments__judge', 'files'))
    items = list(base.order_by('pk'))

    _sheet(wb, 'الطلبات', ['رقم الطلب', 'المدرسة', 'المحافظة', 'المديرية', 'ضابط الارتباط', 'الهاتف', 'البريد', 'المجال',
                          'المسار', 'المشروع', 'المشرف', 'الطلبة', 'الحالة', 'مرحلة التحكيم', 'تاريخ الإرسال', 'الإقرارات',
                          'التحكيم', 'المتوسط', 'فرق الدرجات', 'عدد الملفات', 'تجريبي'],
           [[s.ref, s.school_name, str(getattr(getattr(s.owner, 'profile', None), 'governorate', '') or '') if s.owner_id else '',
             str(getattr(getattr(s.owner, 'profile', None), 'directorate', '') or '') if s.owner_id else '',
             s.contact_person, s.phone, s.email, str(s.field or ''), str(s.track or ''), s.project_title, s.supervisor,
             ' | '.join(s.team_members.splitlines()), STATUS.get(s.status, s.status), s.get_judging_stage_display(),
             _dt(s.sent_at), 'نعم' if s.declared_at else '', s.judging_progress, s.avg_score, s.score_spread,
             s.files.count(), 'نعم' if s.is_simulation else ''] for s in items])

    rows = []
    for s in items:
        for a in s.assignments.all():
            sc = ' | '.join(f'{x.criterion.name}: {float(x.value):g}/{x.criterion.max_score}' for x in a.scores.all())
            rows.append([s.ref, s.project_title, a.judge.get_full_name() or a.judge.email, a.get_stage_display(),
                         _dt(a.assigned_at), _dt(a.completed_at), a.total if a.is_done else '', a.get_recommendation_display(),
                         sc, a.comment, _dt(a.coi_confirmed_at)])
    _sheet(wb, 'التحكيم', ['رقم الطلب', 'المشروع', 'المحكّم', 'المرحلة', 'أُسند', 'انتهى', 'الدرجة /100', 'التوصية',
                          'الدرجات', 'ملاحظات', 'إقرار عدم التعارض'], rows)

    ranked = []
    groups = {}
    for s in items:
        if s.status in ('judging', 'accepted', 'winner', 'rejected') and s.stage_done():
            groups.setdefault(str(s.track or 'بدون مسار'), []).append(s)
    for track, subs_ in groups.items():
        for r in workflow.rank_rows(subs_, cycle.tie_break if cycle else 'criterion'):
            ranked.append([track, r['suggest'], r['s'].ref, r['s'].school_name, r['s'].project_title, r['avg'],
                           'تعادل' if r['tie'] else '', r['s'].score_spread, STATUS.get(r['s'].status)])
    _sheet(wb, 'الترتيب', ['المسار', 'الترتيب', 'رقم الطلب', 'المدرسة', 'المشروع', 'المتوسط', 'تعادل؟', 'فرق الدرجات', 'الحالة'], ranked)

    ws = Winner.objects.filter(cycle=cycle) if cycle else Winner.objects.all()
    _sheet(wb, 'الفائزون', ['المركز', 'الفئة', 'المدرسة', 'المشروع', 'رقم الطلب'],
           [[w.rank, str(w.category or ''), w.school_name, w.project_title, w.submission.ref if w.submission_id else '']
            for w in ws.select_related('category', 'submission').order_by('rank')])

    aps = Appeal.objects.filter(submission__in=items).select_related('submission')
    _sheet(wb, 'الاعتراضات', ['رقم الطلب', 'المدرسة', 'التاريخ', 'الاعتراض', 'الحالة', 'الرد'],
           [[a.submission.ref, a.submission.school_name, _dt(a.created_at), a.body, a.get_status_display(), a.response] for a in aps])

    invs = Invitation.objects.filter(cycle=cycle) if cycle else Invitation.objects.none()
    _sheet(wb, 'الدعوات', ['المدعو', 'البريد', 'الرد', 'عدد الحضور', 'ملاحظة', 'وقت الرد'],
           [[i.name, i.email, i.get_rsvp_display(), i.guests, i.note, _dt(i.responded_at)] for i in invs])

    owners = {s.owner_id for s in items if s.owner_id}
    _sheet(wb, 'المدارس', ['المدرسة', 'ضابط الارتباط', 'البريد', 'الهاتف', 'المحافظة', 'المديرية', 'مفعّل', 'تاريخ التسجيل'],
           [[p.school_name, p.contact_person, p.user.email, p.phone, str(p.governorate or ''), str(p.directorate or ''),
             'نعم' if p.user.is_active else 'لا', _dt(p.created_at)]
            for p in Profile.objects.filter(user_id__in=owners).select_related('user', 'governorate', 'directorate')])

    js = {}
    for a in Assignment.objects.filter(submission__in=items).select_related('judge'):
        d = js.setdefault(a.judge_id, [a.judge.get_full_name() or a.judge.email, a.judge.email, 0, 0])
        d[2] += 1
        d[3] += 1 if a.is_done else 0
    _sheet(wb, 'المحكّمون', ['المحكّم', 'البريد', 'المسند', 'المنتهي'], list(js.values()))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def response(cycle, subs=None, name=''):
    data = cycle_workbook(cycle, subs)
    resp = HttpResponse(data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    fname = name or f'iaj-{(cycle.year if cycle else "all")}-{cycle.pk if cycle else "all"}-{timezone.localdate()}.xlsx'
    resp['Content-Disposition'] = f'attachment; filename="{fname}"'
    return resp
